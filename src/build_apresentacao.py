"""build_apresentacao.py — entrega A9 do WP1A: apresentação do projeto.
O conteúdo é declarado uma única vez em SLIDES e renderizado em dois formatos (PPTX editável e PDF),
para que as duas versões nunca divirjam. Saída: apresentacao/APRESENTACAO-WP1A.{pptx,pdf}"""
import os, textwrap
RAIZ = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")
FIG  = os.path.join(RAIZ, "apresentacao", "figuras")
BASE = os.path.join(RAIZ, "apresentacao", "APRESENTACAO-WP1A")

# cada slide: titulo, subtitulo, itens (bullets), figura, rodape
SLIDES = [
 dict(capa=True,
      titulo="Benchmark reproduzível de métodos clássicos\nde aprendizado de máquina para diagnóstico\nde falhas no Tennessee Eastman Process",
      itens=["Anderson Ricardo de Freitas Rocha",
             "Professor responsável: Clarimar José Coelho",
             "MEI0028 – Modelagem e Simulação · WP1A · PUC Goiás · 2026"]),
 dict(titulo="O problema", subtitulo="Diagnosticar falhas em uma planta química a partir dos sensores",
      itens=["Uma planta química é monitorada por dezenas de variáveis acopladas.",
             "Quando algo sai do normal, é preciso dizer não só que há falha, mas qual é.",
             "Formulado como classificação: cada leitura dos sensores vai para uma de 21 classes.",
             "O erro custa caro: parada de produção, perda de produto, risco de segurança."]),
 dict(titulo="De onde vem o problema", subtitulo="Trinta anos de uma referência que continua em uso",
      itens=["1993 — Downs e Vogel publicam o Tennessee Eastman Process, simulador de uma planta real.",
             "Vira o ambiente padrão para comparar métodos de diagnóstico de falhas.",
             "2012 — Yin et al. fazem a comparação clássica de métodos estatísticos e de aprendizado.",
             "2017 — Rieth et al. publicam uma versão estendida com 500 execuções por condição.",
             "Hoje — a literatura recente usa quase só arquiteturas profundas, com acurácias acima de 99%."]),
 dict(titulo="A base de dados", subtitulo="Rieth et al. (2017) · Harvard Dataverse · domínio público",
      itens=["52 variáveis: 41 medidas de sensores e 11 comandos de válvula, a cada 3 minutos.",
             "21 condições: operação normal e 20 modos de falha.",
             "500 execuções independentes por condição, cada uma com semente própria.",
             "15,33 milhões de linhas ao todo; 10,08 milhões só no conjunto de teste.",
             "A falha entra depois do início da execução: 1 hora no treino, 8 horas no teste."]),
 dict(titulo="O que a literatura recente não faz", subtitulo="Treze trabalhos de 2019 a 2026, analisados em doze parâmetros",
      itens=["Só 4 dos 13 declaram dividir os dados por execução completa.",
             "Dividir por amostra permite ao modelo reconhecer vizinhos já vistos: acurácia inflada.",
             "Nenhum reporta o MCC, métrica menos sensível ao desbalanceamento entre classes.",
             "Nenhum mede o custo computacional de todas as famílias no mesmo equipamento.",
             "Só 3 declaram o que fazem com as amostras anteriores à falha.",
             "Quem divide por execução usa boosting e não publica o resultado dos métodos clássicos."]),
 dict(titulo="Como fizemos", subtitulo="Um protocolo único, declarado por inteiro",
      figura="fig00_fluxograma_protocolo.png", fig_larg=0.32,
      rodape="A execução completa é a unidade e nunca se parte entre treino e teste. Hiperparâmetros escolhidos "
             "só na validação. O teste é lido uma única vez, depois de tudo congelado."),
 dict(titulo="Os seis modelos comparados", subtitulo="Todos sob o mesmo protocolo, os mesmos dados e o mesmo computador",
      itens=["Regressão logística — a fronteira linear, referência mais simples.",
             "Árvore de decisão — uma sequência de perguntas sobre as variáveis.",
             "Random Forest — centenas de árvores votando.",
             "Gradient boosting — árvores que corrigem o erro das anteriores.",
             "XGBoost — a versão otimizada do boosting, hoje dominante em dados tabulares.",
             "SVM aproximado — margem máxima, com o núcleo aproximado para caber na escala."]),
 dict(titulo="Resultados", subtitulo="10,08 milhões de amostras de 10.500 execuções nunca vistas",
      figura="fig01_f1_mcc_comparativo.png", fig_larg=0.78,
      rodape="XGBoost lidera (F1 macro 0,817), seguido de Random Forest (0,794) e da árvore isolada (0,741). "
             "A ordem entre os modelos se mantém quando se troca o conjunto de execuções avaliadas."),
 dict(titulo="Três falhas que ninguém separa", subtitulo="Falhas 3, 9 e 15 confundem-se com a operação normal",
      figura="fig03_matriz_confusao_xgboost.png", fig_larg=0.36,
      rodape="Fora desse bloco a matriz é quase diagonal. O mesmo padrão aparece em trabalhos independentes "
             "sobre a mesma base, o que sugere um limite dos dados e não do método."),
 dict(titulo="Desempenho custa tempo", subtitulo="A pergunta que a literatura não responde",
      figura="fig05_tradeoff_f1_vs_inferencia.png", fig_larg=0.60,
      rodape="A árvore de decisão entrega 91% do F1 macro do XGBoost com inferência 194 vezes mais rápida "
             "e modelo sete vezes menor."),
 dict(titulo="O que fica",
      itens=["Uma linha de base verificável: protocolo, manifesto de divisão e predições publicados.",
             "Dividir por amostra em vez de execução infla o F1 macro entre 0,013 e 0,029.",
             "O rótulo dado às amostras anteriores à falha muda o resultado em até nove pontos.",
             "Métodos clássicos ainda são competitivos quando o custo entra na comparação.",
             "Repositório público: github.com/ricardoarfr/MEI0028-modelagem-e-simulacao"]),
]
ESCURO, MEIO, REALCE = "#1A1A1A", "#555555", "#1F4E79"

def gerar_pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    rgb = lambda h: RGBColor(int(h[1:3],16), int(h[3:5],16), int(h[5:7],16))
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    L, LARG = Inches(0.9), Inches(11.5)
    def caixa(s, txt, esq, topo, larg, tam, cor, negrito=False, espaco=8):
        tf = s.shapes.add_textbox(esq, topo, larg, Inches(0.6)).text_frame; tf.word_wrap = True
        for i, linha in enumerate(txt if isinstance(txt, list) else txt.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT; p.space_after = Pt(espaco)
            r = p.add_run(); r.text = linha
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(tam), negrito, rgb(cor), "Arial"
    for d in SLIDES:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        if d.get("capa"):
            caixa(s, d["titulo"], L, Inches(1.9), LARG, 34, REALCE, True)
            caixa(s, d["itens"], L, Inches(4.5), LARG, 18, MEIO, espaco=10); continue
        caixa(s, d["titulo"], L, Inches(0.6), LARG, 30, REALCE, True)
        if d.get("subtitulo"): caixa(s, d["subtitulo"], L, Inches(1.45), LARG, 16, MEIO)
        if d.get("itens"): caixa(s, ["•  " + i for i in d["itens"]], L, Inches(2.1), LARG, 19, ESCURO, espaco=14)
        if d.get("figura"):
            cam = os.path.join(FIG, d["figura"])
            if os.path.exists(cam):
                larg = Inches(13.333 * d["fig_larg"])
                s.shapes.add_picture(cam, Inches((13.333 - 13.333*d["fig_larg"]) / 2), Inches(1.5), width=larg)
        if d.get("rodape"): caixa(s, d["rodape"], L, Inches(6.4), LARG, 16, MEIO)
    prs.save(BASE + ".pptx"); return len(SLIDES)

def gerar_pdf():
    import matplotlib; matplotlib.use("Agg")
    import matplotlib.pyplot as plt, matplotlib.image as mpimg
    from matplotlib.backends.backend_pdf import PdfPages
    P = lambda pt: pt / 72 / 7.5            # pontos → fração da altura da página
    with PdfPages(BASE + ".pdf") as pdf:
        for d in SLIDES:
            fig = plt.figure(figsize=(13.333, 7.5)); fig.patch.set_facecolor("white")
            x = 0.9 / 13.333
            if d.get("capa"):
                y = 1 - 2.6/7.5
                for linha in d["titulo"].split("\n"):
                    fig.text(x, y, linha, fontsize=26, color=REALCE, fontweight="bold", va="top"); y -= P(40)
                y = 1 - 4.7/7.5
                for linha in d["itens"]:
                    fig.text(x, y, linha, fontsize=14, color=MEIO, va="top"); y -= P(28)
                pdf.savefig(fig); plt.close(fig); continue
            fig.text(x, 1 - 0.75/7.5, d["titulo"], fontsize=23, color=REALCE, fontweight="bold", va="top")
            if d.get("subtitulo"):
                fig.text(x, 1 - 1.55/7.5, d["subtitulo"], fontsize=13, color=MEIO, va="top")
            if d.get("itens"):
                y = 1 - 2.3/7.5
                for it in d["itens"]:
                    linhas = textwrap.wrap(it, 96)
                    fig.text(x, y, "•  " + linhas[0], fontsize=14.5, color=ESCURO, va="top"); y -= P(24)
                    for extra in linhas[1:]:
                        fig.text(x + 0.016, y, extra, fontsize=14.5, color=ESCURO, va="top"); y -= P(24)
                    y -= P(12)
            if d.get("figura"):
                cam = os.path.join(FIG, d["figura"])
                if os.path.exists(cam):
                    im = mpimg.imread(cam); h_im, w_im = im.shape[:2]
                    larg = d["fig_larg"]; alt = larg * 13.333 / 7.5 * h_im / w_im
                    topo, disp = 1 - 1.95/7.5, (6.25 - 1.95) / 7.5
                    if alt > disp: larg *= disp / alt; alt = disp
                    ax = fig.add_axes([(1 - larg) / 2, topo - alt, larg, alt]); ax.imshow(im); ax.axis("off")
            if d.get("rodape"):
                y = 1 - 6.55/7.5
                for linha in textwrap.wrap(d["rodape"], 118):
                    fig.text(x, y, linha, fontsize=12.5, color=MEIO, va="top"); y -= P(20)
            pdf.savefig(fig); plt.close(fig)

n = gerar_pptx(); gerar_pdf()
print(f"gravado {BASE}.pptx e {BASE}.pdf | {n} slides")
